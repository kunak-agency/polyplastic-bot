"""Парсеры для загрузки документов HR: PDF (текстовый слой или OCR), PPTX, XLSX, txt/md.

Каждый парсер возвращает ParsedDocument: чанки в том же формате, что у корпуса, и номера
страниц, которым нельзя доверять без вычитки (OCR с низкой уверенностью или пустая страница).

- PDF: текст берётся из текстового слоя (pypdfium2). Страница почти без текста считается
  сканом и идёт в Tesseract `rus+eng --psm 4` при 300 dpi, если tesseract есть в системе.
  Затем пост-обработка (омоглифы, переносы, колонтитулы, номера пунктов) и обычный
  клаузный чанкинг `chunker.chunk_text`. Если пунктов `N.N` нет, чанк = страница.
- PPTX: слайд = чанк (заголовок, текст фигур, таблицы строками `a | b`, текст SmartArt, заметки
  докладчика).
- XLSX: строка = чанк «Заголовок: значение | …»; первая непустая строка листа — заголовки.
"""

from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from chunker import CLAUSE_RE, SECTION_RE, Chunk, Clause, chunk_text

log = logging.getLogger("rag.parsers")

OCR_LANG = "rus+eng"
OCR_DPI = 300
OCR_PSM = "4"
OCR_TIMEOUT_S = 120
OCR_LOW_CONF = 70.0          # средняя уверенность Tesseract по словам, ниже — страница «под вычитку»
TEXT_LAYER_MIN_CHARS = 30    # меньше непробельных символов в текстовом слое — считаем страницу сканом
SUPPORTED = {".pdf", ".pptx", ".xlsx", ".txt", ".md"}


class UnsupportedDocument(ValueError):
    """Формат не поддерживается (→ 415)."""


class CorruptDocument(ValueError):
    """Формат поддерживается, но файл не открывается или не в той кодировке (→ 422)."""


@dataclass
class ParsedDocument:
    doc: str
    kind: str                                   # pdf | pptx | xlsx | text
    chunks: list[Chunk]
    low_confidence_pages: list[int] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)

    @property
    def clauses(self) -> int:
        return sum(len(c.clauses) for c in self.chunks)


# --- пост-обработка текста (OCR и текстовый слой) ---

_CYR = "а-яёА-ЯЁ"
# латиница, которую OCR путает с кириллицей
_LAT2CYR = str.maketrans({
    "A": "А", "B": "В", "C": "С", "E": "Е", "H": "Н", "K": "К", "M": "М", "O": "О", "P": "Р",
    "T": "Т", "X": "Х", "Y": "У", "a": "а", "c": "с", "e": "е", "o": "о", "p": "р", "x": "х",
    "y": "у", "k": "к", "m": "м", "u": "и", "n": "п", "r": "г",
})
_MIXED_WORD = re.compile(rf"\b(?=\w*[{_CYR}])(?=\w*[A-Za-z])[\w]+\b")
_HYPHEN_BREAK = re.compile(r"([а-яёА-ЯЁA-Za-z])[-‐‑]\s*\n\s*([а-яёa-z])")
_SOFT_HYPHEN = "­"
# «6,1.» / «6,1,3» / «6.1 .» в начале строки → «6.1.». Пробел между цифрами не чиним:
# «3. 15 минут» в заголовке неотличим от «3.15».
_CLAUSE_COMMA = re.compile(r"^(\s*)(\d{1,2}),(\d{1,2})(?:[.,](\d{1,2}))?[.,]?(?=\s)", re.MULTILINE)
_CLAUSE_SPACE_DOT = re.compile(r"^(\s*\d{1,2}\.\d{1,2}(?:\.\d{1,2})?) \.(?=\s)", re.MULTILINE)
_PAGE_NUM_LINE = re.compile(r"^\s*(?:[-–—]\s*)?(?:стр\.?|страница|page)?\s*\d{1,4}\s*(?:из\s*\d{1,4})?(?:\s*[-–—])?\s*$",
                            re.IGNORECASE)


def fix_homoglyphs(text: str) -> str:
    """Латинские буквы внутри кириллических слов → кириллица («пpопуск» → «пропуск»).

    Слова целиком латиницей (MAX, Wi-Fi, e-mail) не трогаем: по одному слову не отличить
    английский термин от ошибки OCR.
    """
    return _MIXED_WORD.sub(lambda m: m.group(0).translate(_LAT2CYR), text)


def join_hyphenation(text: str) -> str:
    """«трудо-\\nвого» → «трудового»; мягкие переносы удаляются."""
    text = text.replace(_SOFT_HYPHEN + "\n", "").replace(_SOFT_HYPHEN, "")
    return _HYPHEN_BREAK.sub(r"\1\2", text)


def fix_clause_numbers(text: str) -> str:
    def repl(m: re.Match) -> str:
        parts = [m.group(2), m.group(3)] + ([m.group(4)] if m.group(4) else [])
        return m.group(1) + ".".join(parts) + "."
    return _CLAUSE_SPACE_DOT.sub(r"\1.", _CLAUSE_COMMA.sub(repl, text))


def _norm_furniture(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def _is_furniture(line: str, furniture: set[str]) -> bool:
    stripped = line.strip()
    if CLAUSE_RE.match(stripped) or SECTION_RE.match(stripped):
        return False  # пункт или раздел никогда не считаем колонтитулом
    return _norm_furniture(line) in furniture or bool(_PAGE_NUM_LINE.match(line))


def strip_headers_footers(pages: list[str], edge_lines: int = 2) -> list[str]:
    """Удаляет колонтитулы: строки из первых/последних `edge_lines` строк страницы, которые
    (с точностью до цифр) повторяются на половине страниц и больше, и строки-номера страниц.
    """
    split = [[l for l in p.split("\n")] for p in pages]
    candidates: Counter[str] = Counter()
    for lines in split:
        nonempty = [l for l in lines if l.strip()]
        edges = set(_norm_furniture(l) for l in nonempty[:edge_lines] + nonempty[-edge_lines:])
        candidates.update(e for e in edges if e)
    min_repeats = max(2, (len(pages) + 1) // 2)
    furniture = {k for k, n in candidates.items() if n >= min_repeats} if len(pages) >= 2 else set()

    out = []
    for lines in split:
        nonempty_idx = [i for i, l in enumerate(lines) if l.strip()]
        edge_idx = set(nonempty_idx[:edge_lines] + nonempty_idx[-edge_lines:])
        kept = [l for i, l in enumerate(lines)
                if not (i in edge_idx and _is_furniture(l, furniture))]
        out.append("\n".join(kept))
    return peel_header_blocks(out)


_WORD = re.compile(rf"[{_CYR}A-Za-z]{{3,}}")
HEADER_BLOCK_LINES = 8      # шапка скана (штамп СТО: организация, шифр, «Лист N», «Листов M», название)
HEADER_LINE_MAX_WORDS = 8


def _header_words(line: str) -> list[str]:
    return [w.lower() for w in _WORD.findall(line)]


def peel_header_blocks(pages: list[str], block_lines: int = HEADER_BLOCK_LINES) -> list[str]:
    """Срезает многострочную шапку-штамп в начале страницы скана.

    У сканов по СТО шапка — 5–8 коротких строк («ООО …», шифр «П 3.11-07», «Лист 7», «Листов 23»,
    название документа), и OCR каждый раз читает её по-разному («II 3.11-07», «13,1107 [Г положниЕ»),
    поэтому точное совпадение строк из `strip_headers_footers` её не ловит. Здесь словарь шапки —
    слова, которые встречаются в первых `block_lines` строках на половине страниц и больше. На каждой
    странице находим последнюю «шапочную» строку в верхнем блоке (короткая, ≥ половины слов из
    словаря) и срезаем всё до неё включительно: мусор OCR между строками шапки уходит вместе с ней.
    Блок кончается на первом пункте или разделе — их не трогаем никогда.
    """
    if len(pages) < 3:
        return pages
    split = [p.split("\n") for p in pages]
    tops: list[list[tuple[int, str]]] = []
    df: Counter[str] = Counter()
    for lines in split:
        top: list[tuple[int, str]] = []
        for i, l in enumerate(lines):
            s = l.strip()
            if not s:
                continue
            if CLAUSE_RE.match(s) or SECTION_RE.match(s) or len(top) >= block_lines:
                break
            top.append((i, s))
        tops.append(top)
        df.update(set(w for _, s in top for w in _header_words(s)))
    vocab = {w for w, n in df.items() if n >= max(3, (len(pages) + 1) // 2)}
    if not vocab:
        return pages
    out = []
    for lines, top in zip(split, tops):
        cut = -1
        for i, s in top:
            words = _header_words(s)
            if words and len(s.split()) <= HEADER_LINE_MAX_WORDS and \
                    sum(w in vocab for w in words) * 2 >= len(words):
                cut = i
        out.append("\n".join(lines[cut + 1:]) if cut >= 0 else "\n".join(lines))
    return out


def postprocess_pages(pages: list[str]) -> list[str]:
    pages = strip_headers_footers([p.replace("\r\n", "\n").replace("\r", "\n") for p in pages])
    return [fix_clause_numbers(fix_homoglyphs(join_hyphenation(p))) for p in pages]


# --- OCR ---

def tesseract_available(lang: str = OCR_LANG) -> bool:
    exe = shutil.which("tesseract")
    if not exe:
        return False
    try:
        out = subprocess.run([exe, "--list-langs"], capture_output=True, text=True, timeout=10)
    except (OSError, subprocess.SubprocessError):
        return False
    langs = set((out.stdout + out.stderr).split())
    return all(l in langs for l in lang.split("+"))


def ocr_image(png: bytes, lang: str = OCR_LANG) -> tuple[str, float]:
    """PNG → (текст, средняя уверенность слов 0–100). Tesseract вызывается без shell."""
    exe = shutil.which("tesseract")
    if not exe:
        raise RuntimeError("tesseract не установлен")
    res = subprocess.run([exe, "stdin", "stdout", "-l", lang, "--psm", OCR_PSM, "tsv"],
                         input=png, capture_output=True, timeout=OCR_TIMEOUT_S, check=True)
    rows = res.stdout.decode("utf-8", errors="replace").splitlines()
    lines: dict[tuple[int, int, int], list[str]] = {}
    confs: list[float] = []
    for row in rows[1:]:
        cols = row.split("\t")
        if len(cols) < 12 or cols[0] != "5":  # level 5 = слово
            continue
        word = cols[11].strip()
        if not word:
            continue
        try:
            conf = float(cols[10])
        except ValueError:
            conf = -1
        if conf >= 0:
            confs.append(conf)
        key = (int(cols[2]), int(cols[3]), int(cols[4]))  # block, paragraph, line
        lines.setdefault(key, []).append(word)
    text = "\n".join(" ".join(ws) for _, ws in sorted(lines.items()))
    return text, (sum(confs) / len(confs) if confs else 0.0)


# --- PDF ---

def _page_marked(pages: list[str]) -> str:
    return "\n".join(f"--- стр. {i} ---\n{t}" for i, t in enumerate(pages, start=1))


@dataclass
class PdfText:
    pages: list[str]                 # текст страниц после пост-обработки
    low_confidence_pages: list[int]
    ocr_pages: list[int]
    ocr_conf: dict[int, float]       # страница → средняя уверенность Tesseract по словам


def pdf_text(data: bytes, doc: str, use_ocr: bool | None = None) -> PdfText:
    """PDF → постраничный текст: текстовый слой или OCR, затем пост-обработка."""
    import pypdfium2 as pdfium  # noqa: PLC0415

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise CorruptDocument("не удалось открыть PDF") from exc
    ocr_ok = tesseract_available() if use_ocr is None else use_ocr
    raw_pages: list[str] = []
    low_conf: list[int] = []
    ocr_pages: list[int] = []
    confs: dict[int, float] = {}
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            text = textpage.get_text_range().replace("\r\n", "\n")
            textpage.close()
            if len(re.sub(r"\s", "", text)) < TEXT_LAYER_MIN_CHARS:
                if ocr_ok:
                    bitmap = page.render(scale=OCR_DPI / 72, grayscale=True)
                    buf = io.BytesIO()
                    bitmap.to_pil().save(buf, format="PNG")
                    text, conf = ocr_image(buf.getvalue())
                    ocr_pages.append(i + 1)
                    confs[i + 1] = conf
                    if conf < OCR_LOW_CONF or not text.strip():
                        low_conf.append(i + 1)
                else:
                    low_conf.append(i + 1)  # скан без OCR — страницу никто не прочитал
            raw_pages.append(text)
            page.close()
    finally:
        pdf.close()
    if ocr_pages:
        log.info("%s: OCR на %d стр., низкая уверенность: %s", doc, len(ocr_pages), low_conf)
    elif low_conf and not ocr_ok:
        log.warning("%s: %d стр. без текстового слоя, tesseract недоступен", doc, len(low_conf))
    return PdfText(postprocess_pages(raw_pages), low_conf, ocr_pages, confs)


def parse_pdf(data: bytes, doc: str, use_ocr: bool | None = None) -> ParsedDocument:
    t = pdf_text(data, doc, use_ocr)
    pages = t.pages
    chunks = chunk_text(_page_marked(pages), default_doc=doc)
    if not chunks:  # нет нумерованных пунктов — чанк = страница
        chunks = [Chunk(doc=doc, section=None, clauses=[Clause(f"стр. {i}", " ".join(p.split()), i)])
                  for i, p in enumerate(pages, start=1) if p.strip()]
    for ch in chunks:
        ch.doc = doc
    return ParsedDocument(doc=doc, kind="pdf", chunks=chunks, low_confidence_pages=t.low_confidence_pages,
                          ocr_pages=t.ocr_pages)


def export_pdf_markdown(data: bytes, doc: str, source: str | None = None,
                        use_ocr: bool | None = None) -> tuple[str, PdfText]:
    """PDF → markdown корпуса (front-matter + `--- стр. N ---`), чтобы один раз сделать OCR,
    вычитать текст глазами и класть его в RAG_CORPUS_DIR вместо скана."""
    t = pdf_text(data, doc, use_ocr)
    meta = [f"doc: {doc}"] + ([f"source: {source}"] if source else [])
    if t.ocr_pages:
        meta.append(f"ocr: tesseract {OCR_LANG} {OCR_DPI} dpi, стр. {len(t.ocr_pages)}")
    meta.append("low_confidence_pages: " + (", ".join(map(str, t.low_confidence_pages)) or "-"))
    return "---\n" + "\n".join(meta) + "\n---\n" + _page_marked(t.pages) + "\n", t


# --- PPTX ---

def _shape_texts(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        if t.strip():
            out.append(t)
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # группа
        for sub in shape.shapes:
            out.extend(_shape_texts(sub))
    return out


_NS_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _smartart_texts(shape, slide) -> list[str]:
    """Текст SmartArt-диаграммы: python-pptx его не отдаёт, он лежит в части diagramData.

    Абзацы узлов (`dgm:pt`, кроме служебных переходов и узлов представления) по порядку
    документа, по строке на абзац. Битая или отсутствующая часть — пустой список.
    """
    from lxml import etree  # noqa: PLC0415 — зависимость python-pptx

    out: list[str] = []
    for rel in shape._element.iter(f"{{{_NS_DGM}}}relIds"):
        rid = rel.get(f"{{{_NS_R}}}dm")
        try:
            root = etree.fromstring(slide.part.related_part(rid).blob)
        except Exception:  # noqa: BLE001 — нет части или не XML: просто без текста диаграммы
            log.warning("SmartArt без данных: rId=%s", rid)
            continue
        for pt in root.iter(f"{{{_NS_DGM}}}pt"):
            if pt.get("type") in ("parTrans", "sibTrans", "pres"):
                continue
            for para in pt.iter(f"{{{_NS_A}}}p"):
                t = " ".join("".join(para.itertext()).split())
                if t:
                    out.append(t)
    return out


def parse_pptx(data: bytes, doc: str) -> ParsedDocument:
    from pptx import Presentation  # noqa: PLC0415

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # битый zip, не pptx
        raise CorruptDocument("не удалось открыть PPTX") from exc
    chunks: list[Chunk] = []
    empty: list[int] = []
    for n, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = " ".join(title_shape.text.split()) if title_shape is not None and title_shape.has_text_frame else ""
        parts = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            parts.extend(_shape_texts(shape))
            parts.extend(_smartart_texts(shape, slide))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            if notes.strip():
                parts.append(f"Заметки: {notes.strip()}")
        body = "\n".join(([title] if title else []) + parts).strip()
        if not body:
            empty.append(n)  # слайд-картинка: текста нет, OCR слайдов не делаем
            continue
        chunks.append(Chunk(doc=doc, section=title or None,
                            clauses=[Clause(f"слайд {n}", " ".join(body.split()), n)]))
    return ParsedDocument(doc=doc, kind="pptx", chunks=chunks, low_confidence_pages=empty)


# --- XLSX ---

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def parse_xlsx(data: bytes, doc: str) -> ParsedDocument:
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptDocument("не удалось открыть XLSX") from exc
    chunks: list[Chunk] = []
    try:
        sheets = wb.worksheets
        for ws in sheets:
            sheet_doc = doc if len(sheets) == 1 else f"{doc} / {ws.title}"
            header: list[str] | None = None
            for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [_cell(v) for v in row]
                if not any(cells):
                    continue
                if header is None:
                    header = [c or f"Колонка {i}" for i, c in enumerate(cells, start=1)]
                    continue
                pairs = [f"{header[i] if i < len(header) else f'Колонка {i + 1}'}: {c}"
                         for i, c in enumerate(cells) if c]
                chunks.append(Chunk(doc=sheet_doc, section=ws.title,
                                    clauses=[Clause(f"строка {r}", " | ".join(pairs), None)]))
    finally:
        wb.close()
    return ParsedDocument(doc=doc, kind="xlsx", chunks=chunks)


# --- диспетчер ---

def parse_document(filename: str, data: bytes, doc: str | None = None) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    name = (doc or "").strip() or Path(filename).stem or "документ"
    if suffix == ".pdf":
        return parse_pdf(data, name)
    if suffix == ".pptx":
        return parse_pptx(data, name)
    if suffix == ".xlsx":
        return parse_xlsx(data, name)
    if suffix in (".txt", ".md"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise CorruptDocument("текст должен быть в UTF-8") from exc
        chunks = chunk_text(text, default_doc=name)
        if doc:
            for ch in chunks:
                ch.doc = name
        return ParsedDocument(doc=chunks[0].doc if chunks else name, kind="text", chunks=chunks)
    raise UnsupportedDocument(f"формат {suffix or '(без расширения)'} не поддерживается; "
                              f"ожидается {', '.join(sorted(SUPPORTED))}")


def _main(argv: list[str] | None = None) -> int:
    """`python parsers.py export файл.pdf --doc ПВТР -o corpus/pvtr.md` — OCR в markdown корпуса.
    В stdout — только сводка по страницам (номер, уверенность), текст документа не печатается."""
    import argparse  # noqa: PLC0415

    ap = argparse.ArgumentParser(prog="python parsers.py")
    sub = ap.add_subparsers(dest="cmd", required=True)
    ex = sub.add_parser("export", help="PDF → markdown корпуса")
    ex.add_argument("pdf", type=Path)
    ex.add_argument("--doc", required=True, help="имя документа для sources[].doc")
    ex.add_argument("-o", "--out", type=Path, required=True)
    args = ap.parse_args(argv)
    md, t = export_pdf_markdown(args.pdf.read_bytes(), args.doc, source=args.pdf.name)
    args.out.write_text(md, encoding="utf-8")
    for n in range(1, len(t.pages) + 1):
        conf = f"{t.ocr_conf[n]:.1f}" if n in t.ocr_conf else "текст"
        print(f"стр. {n}: {conf}{'  LOW' if n in t.low_confidence_pages else ''}")
    chunks = chunk_text(md, default_doc=args.doc)
    print(f"страниц {len(t.pages)}, OCR {len(t.ocr_pages)}, low_confidence {t.low_confidence_pages}, "
          f"чанков {len(chunks)}, пунктов {sum(len(c.clauses) for c in chunks)} → {args.out}")
    return 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    raise SystemExit(_main())
