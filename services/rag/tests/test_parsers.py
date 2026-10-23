"""Парсеры загрузки: PDF (текстовый слой и скан), PPTX, XLSX, пост-обработка OCR.

Файлы-образцы собираются в тесте из синтетического текста, бинарных фикстур в репозитории нет.
PDF с кириллицей требует TTF-шрифт с кириллицей в системе, OCR — tesseract с rus; иначе skip.
"""

import io
from pathlib import Path

import pytest

from parsers import (CorruptDocument, UnsupportedDocument, fix_clause_numbers, fix_homoglyphs, join_hyphenation,
                     parse_document, postprocess_pages, strip_headers_footers, tesseract_available)

FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
]
FONT = next((f for f in FONT_CANDIDATES if Path(f).is_file()), None)
needs_font = pytest.mark.skipif(FONT is None, reason="нет TTF-шрифта с кириллицей")

HEADER = "ООО «Синтетика». Тестовый регламент"
PAGES = [
    ["1. Общие положения",
     "1.1. Регламент синтетический и составлен для тестов загрузки. Сотрудник получает",
     "пропуск в бюро пропусков под подпись в первый рабочий день."],
    ["2. Перерывы",
     "2.1. Перерыв для питания длится тридцать минут и предоставляется в середине смены по",
     "графику, который утверждает мастер участка."],
]


# --- пост-обработка ---

def test_homoglyphs_fixed_only_inside_cyrillic_words():
    assert fix_homoglyphs("Пpи утеpе пpопуска") == "При утере пропуска"
    assert fix_homoglyphs("бот в MAX и Wi-Fi") == "бот в MAX и Wi-Fi"


def test_hyphenation_joined():
    assert join_hyphenation("трудо-\nвого договора") == "трудового договора"
    assert join_hyphenation("сине-\nЗелёный") == "сине-\nЗелёный"  # с заглавной — не перенос


def test_clause_number_ocr_fixes():
    assert fix_clause_numbers("6,1. Текст") == "6.1. Текст"
    assert fix_clause_numbers("6,1,3 Текст") == "6.1.3. Текст"
    assert fix_clause_numbers("6.1 . Текст") == "6.1. Текст"
    assert fix_clause_numbers("3. 15 минут") == "3. 15 минут"


def test_headers_footers_and_page_numbers_removed_but_clauses_kept():
    pages = [f"{HEADER}\n1.{i}. пункт {i}\nтекст\n- {i} -" for i in range(1, 4)]
    out = strip_headers_footers(pages)
    assert all(HEADER not in p and "- " not in p for p in out)
    assert [p.split("\n")[0] for p in out] == ["1.1. пункт 1", "1.2. пункт 2", "1.3. пункт 3"]


def test_postprocess_pipeline():
    out = postprocess_pages(["Шапка\n6,1. Заработная пла-\nта выплачивается двa раза\n1",
                             "Шапка\n6,2. Аванс\n2"])
    assert out[0].strip() == "6.1. Заработная плата выплачивается два раза"


# --- PDF ---

def _text_pdf(pages: list[list[str]]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(False)
    pdf.add_font("cyr", fname=FONT)
    pdf.set_font("cyr", size=11)
    for n, lines in enumerate(pages, start=1):
        pdf.add_page()
        pdf.cell(0, 8, HEADER, new_x="LMARGIN", new_y="NEXT")
        for line in lines:
            pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
        pdf.set_y(-20)
        pdf.cell(0, 8, f"Страница {n} из {len(pages)}")
    return bytes(pdf.output())


def _scan_pdf(pages: list[list[str]]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(FONT, 40)
    images = []
    for lines in pages:
        img = Image.new("L", (2480, 3508), 255)  # A4 при 300 dpi
        d = ImageDraw.Draw(img)
        y = 200
        for line in lines:
            d.text((150, y), line, fill=0, font=font)
            y += 70
        images.append(img.convert("RGB"))
    buf = io.BytesIO()
    images[0].save(buf, format="PDF", save_all=True, append_images=images[1:], resolution=300)
    return buf.getvalue()


@needs_font
def test_text_pdf_clauses_pages_and_header_removed():
    parsed = parse_document("reg.pdf", _text_pdf(PAGES), doc="Тестовый регламент")
    assert parsed.kind == "pdf" and parsed.low_confidence_pages == [] and parsed.ocr_pages == []
    clauses = {cl.clause_id: cl for ch in parsed.chunks for cl in ch.clauses}
    assert set(clauses) == {"1.1", "2.1"}
    assert (clauses["1.1"].page, clauses["2.1"].page) == (1, 2)
    assert "пропуск в бюро пропусков" in clauses["1.1"].text
    assert all("Синтетика" not in ch.body and "Страница" not in ch.body for ch in parsed.chunks)
    assert all(ch.doc == "Тестовый регламент" for ch in parsed.chunks)


@needs_font
def test_scan_pdf_without_tesseract_marks_pages_low_confidence(monkeypatch):
    import parsers
    monkeypatch.setattr(parsers, "tesseract_available", lambda lang=parsers.OCR_LANG: False)
    parsed = parse_document("scan.pdf", _scan_pdf(PAGES))
    assert parsed.chunks == [] and parsed.low_confidence_pages == [1, 2]


@needs_font
@pytest.mark.skipif(not tesseract_available(), reason="tesseract rus+eng не установлен")
def test_scan_pdf_ocr_rus():
    parsed = parse_document("scan.pdf", _scan_pdf(PAGES), doc="Скан")
    assert parsed.ocr_pages == [1, 2]
    clauses = {cl.clause_id: cl for ch in parsed.chunks for cl in ch.clauses}
    assert {"1.1", "2.1"} <= set(clauses)
    assert clauses["2.1"].page == 2
    assert "перерыв для питания" in clauses["2.1"].text.lower()


def test_broken_pdf_rejected():
    with pytest.raises(CorruptDocument):
        parse_document("x.pdf", b"%PDF-1.4 not really")


# --- PPTX / XLSX / прочее ---

def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Первый день"
    s1.placeholders[1].text = "Приходите к 8:00 на пост охраны, там выдадут временный пропуск"
    s1.notes_slide.notes_text_frame.text = "Напомнить про СИЗ"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Контакты"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Роль", "Кабинет"
    table.cell(1, 0).text, table.cell(1, 1).text = "Наставник", "204"
    prs.slides.add_slide(prs.slide_layouts[6])  # пустой слайд
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_slide_is_chunk():
    parsed = parse_document("welcome.pptx", _pptx())
    assert parsed.kind == "pptx" and parsed.doc == "welcome"
    assert [ch.clause_id for ch in parsed.chunks] == ["слайд 1", "слайд 2"]
    first, second = parsed.chunks
    assert first.section == "Первый день" and first.page == 1
    assert "временный пропуск" in first.body and "Заметки: Напомнить про СИЗ" in first.body
    assert "Наставник | 204" in second.body
    assert parsed.low_confidence_pages == [3]  # пустой слайд — без текста


def _xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Сценарий"
    ws.append(["Этап", "День", "Площадка", "Ответственный", "Действие"])
    ws.append(["До выхода", -3, "GTP", "Рекрутер", "Позвонить и напомнить про документы"])
    ws.append([None, None, None, None, None])
    ws.append(["Первый день", 0, "GTP", "Наставник", "Провести по цеху"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_row_is_chunk_with_headers():
    parsed = parse_document("scenario.xlsx", _xlsx(), doc="Сценарий GTP")
    assert parsed.kind == "xlsx"
    assert [ch.clause_id for ch in parsed.chunks] == ["строка 2", "строка 4"]
    assert parsed.chunks[1].clauses[0].text == \
        "Этап: Первый день | День: 0 | Площадка: GTP | Ответственный: Наставник | Действие: Провести по цеху"
    assert parsed.chunks[0].doc == "Сценарий GTP"


def test_markdown_upload_uses_clause_chunker():
    parsed = parse_document("a.md", "1. Раздел\n1.1. Пункт про обед в столовой\n".encode(), doc="Памятка")
    assert parsed.kind == "text" and parsed.chunks[0].doc == "Памятка" and parsed.clauses == 1


def test_unsupported_extension():
    with pytest.raises(UnsupportedDocument):
        parse_document("a.docx", b"PK")
